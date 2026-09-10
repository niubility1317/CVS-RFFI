"""Reproduce method-level association and class-level descriptive diagnostics."""
import csv, gzip, itertools, json, math
from pathlib import Path
import numpy as np
from scipy.stats import pearsonr, spearmanr, rankdata
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

BASE=Path(__file__).parent
OUT=BASE/'method_class_analysis_20260910'
OUT.mkdir(exist_ok=True)
target=json.loads((BASE/'completed_experiments_20260910_2122.json').read_text(encoding='utf-8'))
source=json.loads((BASE/'source_classes_results.json').read_text(encoding='utf-8'))
with gzip.open(OUT/'full_log_audit.json.gz','rt',encoding='utf-8') as stream:
    audit=json.load(stream)
NAMES=list(target['rows'])
SCENES=['leo_clear_weak','leo_low_elev_weak','leo_rain_weak']
methods=[]; classes=[]; scenes=[]; curves=[]; activation=[]
for name in NAMES:
    r=target['rows'][name]; ss=source['rows'][name]['scenarios']
    assert r['checkpoint_epoch']==200 and r['e200_snapshot_model_equal_final']
    assert source['rows'][name]['model_unchanged']
    assert len(r['curve'])==200 and len(audit['rows'][name]['records'])==200
    assert [x['epoch'] for x in r['scores']]==list(range(80,201,10))
    ts=r['scores'][-1]['score']['metrics']
    m=dict(method=name,source_clean=r['final_source_clean'],source_leo=r['final_source_leo'],
        target_clean=ts['clean']['accuracy']*100,target_leo=np.mean([ts[k]['accuracy']*100 for k in SCENES]),
        epoch_hours=r['epoch_hours'],target_test_hours=r['target_test_hours'])
    for c in range(6):
        cs=str(c)
        sc=ss['clean']['per_class'][cs]['accuracy_pct']
        sl=np.mean([ss[k]['per_class'][cs]['accuracy_pct'] for k in SCENES])
        tc=ts['clean']['per_class_accuracy'][cs]*100
        tl=np.mean([ts[k]['per_class_accuracy'][cs]*100 for k in SCENES])
        classes.append(dict(method=name,class_index=c,source_clean=sc,source_leo=sl,target_clean=tc,target_leo=tl,
            clean_domain_gap=sc-tc,source_leo_drop=sc-sl,target_leo_drop=tc-tl,
            interaction=(tc-tl)-(sc-sl),leo_domain_gap=sl-tl))
        for k in ['clean']+SCENES:
            v=ss[k]['per_class'][cs]
            assert v['total']==4500 and math.isclose(v['accuracy_pct'],100*v['correct']/4500,abs_tol=1e-10)
            scenes.append(dict(method=name,class_index=c,scene=k,source_correct=v['correct'],source_total=v['total'],
                source_accuracy=v['accuracy_pct'],target_accuracy=ts[k]['per_class_accuracy'][cs]*100))
    cr=[x for x in classes if x['method']==name]
    m.update(target_min_class=min(x['target_leo'] for x in cr),
        target_bottom2=np.mean(sorted(x['target_leo'] for x in cr)[:2]),
        target_class_std=np.std([x['target_leo'] for x in cr]),
        source_min_class=min(x['source_leo'] for x in cr))
    assert math.isclose(np.mean([x['target_leo'] for x in cr]),m['target_leo'],abs_tol=1e-10)
    methods.append(m)
    byepoch={x['epoch']:x for x in r['curve']}
    for score in r['scores']:
        e=score['epoch']; sm=score['score']['metrics']
        v=dict(method=name,epoch=e,source_clean=byepoch[e]['val_tx_acc'],source_leo=byepoch[e]['stage_source_val_sat_mean_tx'],
            target_leo=np.mean([sm[k]['accuracy']*100 for k in SCENES]))
        for c in range(6):v['target_class'+str(c)]=np.mean([sm[k]['per_class_accuracy'][str(c)]*100 for k in SCENES])
        curves.append(v)
    logs=audit['rows'][name]['records']
    ac=dict(method=name,first_sat_ce_epoch=next((x['epoch'] for x in logs if x.get('train_w_loss_sat_cls_labeled',0)>0),None),
        first_daot_epoch=next((x['epoch'] for x in logs if x.get('train_loss_daot_total',0)>0),None),
        grad_skip_epochs=sum(x.get('train_skipped_nonfinite_grad',0)>0 for x in logs),
        ecrs_successful_steps_max=max(x.get('train_ecrs_cross_rx_successful_steps',0) for x in logs))
    for k in ['lr','train_daot_teacher_view_count','train_loss_daot_total','train_daot_consensus_rate',
        'train_rc4_effective_weighted_coverage','train_pseudo_truth_available','train_w_loss_sat_cls_labeled']:
        ac[k]=logs[-1].get(k)
    activation.append(ac)
for m in methods:
    m['source_leo_delta_vs_B0']=m['source_leo']-methods[0]['source_leo']
    m['target_leo_delta_vs_B0']=m['target_leo']-methods[0]['target_leo']
def correlation(rows,sourcekey='source_leo',targetkey='target_leo'):
    x=np.array([r[sourcekey] for r in rows]); y=np.array([r[targetkey] for r in rows])
    return dict(n=len(rows),pearson=float(pearsonr(x,y).statistic),spearman=float(spearmanr(x,y).statistic))
corr=[]
for label,excluded in [('all',[]),('without_D0',['D0_NO_ORBIT']),('without_B1',['B1_TAIL_LR']),('without_D0_B1',['D0_NO_ORBIT','B1_TAIL_LR'])]:
    corr.append(dict(scope=label,**correlation([r for r in methods if r['method'] not in excluded])))
for name in NAMES:corr.append(dict(scope='leave_out_'+name,**correlation([r for r in methods if r['method']!=name])))
corr.append(dict(scope='source_clean_to_target_leo',**correlation(methods,'source_clean')))
for c in range(6):corr.append(dict(scope='same_class_across_methods_'+str(c),**correlation([r for r in classes if r['class_index']==c])))
for k in SCENES:
    pairs=[]
    for name in NAMES:
        p=[x for x in scenes if x['method']==name and x['scene']==k]
        pairs.append(dict(source_leo=np.mean([x['source_accuracy'] for x in p]),target_leo=np.mean([x['target_accuracy'] for x in p])))
    corr.append(dict(scope='scene_'+k,**correlation(pairs)))
b1=[r for r in classes if r['method']=='B1_TAIL_LR'];b0=[r for r in classes if r['method']=='B0_FIXED']
deltas=[dict(class_index=x['class_index'],source_delta=x['source_leo']-y['source_leo'],target_delta=x['target_leo']-y['target_leo']) for x,y in zip(b1,b0)]

tables={'methods':methods,'classes':classes,'scene_classes':scenes,'correlations':corr,'fixed_epoch_curves':curves,'activation':activation,'B1_B0_class_delta':deltas}
for name,rows in tables.items():
    with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
(OUT/'statistics.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

plt.rcParams.update({'font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.facecolor':'white'})
fig,axes=plt.subplots(1,2,figsize=(12.8,4.8),layout='constrained')
for ax,selected,title in [(axes[0],methods,'All 8 methods: Pearson r = 0.931'),(axes[1],[m for m in methods if m['method'] not in ['B1_TAIL_LR','D0_NO_ORBIT']],'Six nearby methods: rank correlation = 0.143')]:
    for m in selected:
        color='#b04c36' if m['method']=='D0_NO_ORBIT' else '#24658c' if m['method']=='B1_TAIL_LR' else '#444444'
        ax.scatter(m['source_leo'],m['target_leo'],s=55,c=color)
        offsets={'B0_FIXED':(6,-15),'B2_RC4_WEIGHT':(6,-15),'D3_TANGENT':(6,5),'D1_THREE_VIEW':(6,5),'F3_R3_SWAP':(6,-12),'D2_PHYSICAL_ORBIT':(6,5),'B1_TAIL_LR':(-60,9),'D0_NO_ORBIT':(7,5)}
        if ax is axes[1] or m['method'] in ['B0_FIXED','B1_TAIL_LR','D0_NO_ORBIT']:
            ax.annotate(m['method'].split('_')[0],(m['source_leo'],m['target_leo']),xytext=offsets[m['method']],textcoords='offset points')
    ax.set(xlabel='Source V LEO mean accuracy (%)',ylabel='Target LEO mean accuracy (%)',title=title)
    ax.grid(alpha=.2);ax.margins(x=.18,y=.2)
fig.savefig(OUT/'method_correlation.png',dpi=180);plt.close(fig)
fig,axes=plt.subplots(1,2,figsize=(12,5.8),layout='constrained')
for ax,key,title in zip(axes,['source_leo','target_leo'],['Source V LEO gain vs B0 (pp)','Target LEO gain vs B0 (pp)']):
    matrix=np.array([[next(r[key] for r in classes if r['method']==n and r['class_index']==c)-b0[c][key] for c in range(6)] for n in NAMES])
    im=ax.imshow(matrix,cmap='RdBu',vmin=-8,vmax=8,aspect='auto')
    ax.set(xticks=range(6),xticklabels=[f'Class {c}' for c in range(6)],yticks=range(8),yticklabels=NAMES,title=title)
    for (i,j),v in np.ndenumerate(matrix):ax.text(j,i,f'{v:+.2f}',ha='center',va='center',color='white' if abs(v)>5 else 'black',fontsize=9)
fig.colorbar(im,ax=axes,label='Accuracy difference (percentage points)',shrink=.8)
fig.savefig(OUT/'class_gains.png',dpi=180);plt.close(fig)
fig,axes=plt.subplots(1,2,figsize=(12,4.6),layout='constrained')
x=np.arange(6)
for key,label,color in [('source_clean','Source clean','#222222'),('source_leo','Source LEO','#6089a6'),('target_clean','Target clean','#a86949'),('target_leo','Target LEO','#b52e38')]:
    axes[0].plot(x,[v[key] for v in b1],marker='o',label=label,c=color)
axes[0].set(xticks=x,xlabel='Class index',ylabel='Accuracy (%)',ylim=(25,103),title='B1: four observation conditions');axes[0].legend(fontsize=9);axes[0].grid(alpha=.2)
cc=[v for v in curves if v['method']=='B1_TAIL_LR']
for key,label in [('source_leo','Source LEO mean'),('target_leo','Target LEO mean'),('target_class1','Target class 1'),('target_class3','Target class 3')]:
    axes[1].plot([v['epoch'] for v in cc],[v[key] for v in cc],marker='.',label=label)
axes[1].set(xlabel='Fixed epoch',ylabel='Accuracy (%)',title='B1: stored fixed-epoch evaluations');axes[1].legend(fontsize=8,loc='upper center',bbox_to_anchor=(.5,-.16),ncol=2);axes[1].grid(alpha=.2)
fig.savefig(OUT/'B1_classes_and_curve.png',dpi=180);plt.close(fig)

wb=Workbook();wb.remove(wb.active)
for name,rows in tables.items():
    ws=wb.create_sheet(name[:31]);ws.append(list(rows[0]));[ws.append(list(r.values())) for r in rows]
ws=wb.create_sheet('Formula_check');ws.append(['Scope','Pearson formula','Spearman note'])
ws.append(['8 methods','=CORREL(methods!C2:C9,methods!E2:E9)','Rank coefficients recomputed in correlations sheet; descriptive only'])
ws=wb.create_sheet('Sources');ws.append(['Source','Location','Meaning'])
for v in [('Final E200 metrics',str(BASE/'completed_experiments_20260910_2122.json'),'Stored source aggregate and sealed target scores; exploratory'),('Source per-class replay',str(BASE/'source_classes_results.json'),'Frozen E200; 4500 observations per class per scene'),('Full log audit',str(OUT/'full_log_audit.json.gz'),'1600 JSONL records parsed; matching CSV and complete stdout scan'),('Analysis code',str(Path(__file__).resolve()),'All differences in percentage points; no independent-seed inference')]:ws.append(v)
for ws in wb:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='333333')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=min(52,max(15,len(str(col[0].value))+2))
        for cell in col[1:]:
            if isinstance(cell.value,float):cell.number_format='0.0000'
wb.save(OUT/'supporting_data.xlsx')
check=load_workbook(OUT/'supporting_data.xlsx',read_only=True,data_only=False)
assert check['methods'].max_row==9 and check['classes'].max_row==49
assert check['Formula_check']['B2'].value=='=CORREL(methods!C2:C9,methods!E2:E9)'
check.close()
print(json.dumps({'methods':len(methods),'class_rows':len(classes),'scene_class_rows':len(scenes),'fixed_scores':len(curves),'main_correlations':corr[:4],'B1_class_deltas':deltas},ensure_ascii=False))
