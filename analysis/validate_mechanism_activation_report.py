"""Validate numerical claims, spreadsheet closure, links, and delivery mirrors."""
import csv
import gzip
import json
import re
import shutil
from pathlib import Path
from openpyxl import load_workbook

out = Path(__file__).resolve().parent / 'mechanism_activation_audit_20260910'
d = json.load(gzip.open(out / 'evidence.json.gz', 'rt', encoding='utf-8'))
report = (out / 'report.md').read_text(encoding='utf-8')
assert '\ufffd' not in report and not report.startswith('\ufeff')
with (out / 'ecrs_effects.csv').open(encoding='utf-8-sig', newline='') as f:
    effects = {x['row']: x for x in csv.DictReader(f)}
xrows = d['runs']['a1_ecrs_cross_rx_s392005_20260909_r1']['rows']
for n, x in xrows.items():
    m = x['scores'][-1]['score']['metrics']
    leo = [v for k, v in m.items() if k.startswith('leo_')]
    manual = sum(v['correct'] for v in leo) / sum(v['total'] for v in leo) * 100
    assert abs(manual - float(effects[n]['target_leo_mean'])) < 1e-10
    if not n.startswith('X0'):
        assert x['stats']['train_ecrs_cross_rx_identity_grad_norm']['nonzero_epochs'] == 200
        assert x['stats']['train_ecrs_cross_rx_weighted_loss']['nonzero_epochs'] == 200
    if n[:2] in ['X2', 'X3', 'X4', 'X5', 'X6', 'X7']:
        assert x['stats']['train_ecrs_cross_rx_leo_count']['first_nonzero'] == 80
for line in report.splitlines():
    cells = line.strip('|').split('|')
    if cells[0] in effects:
        x = effects[cells[0]]
        for cell, key in zip(cells[1:], ['target_clean', 'target_leo_mean', 'leo_delta_vs_x0_pp', 'leo_delta_vs_x2_pp']):
            assert abs(float(cell) - float(x[key])) <= .000050001
for run in d['runs'].values():
    for x in run['rows'].values():
        assert not x['parse_errors'] and not x['csv_bad_widths']
        assert x['jsonl_records'] == x['csv_records']
        for k, v in x['stats'].items():
            if 'r3_' in k and '_eta_' in k:
                assert v['nonzero_epochs'] == 0
curves = json.loads((out / 'extended_signal_curves.json').read_text(encoding='utf-8'))
x600 = curves['X2_E600']
assert x600[-1]['epoch'] == 184
assert len([x for x in x600 if x['epoch'] >= 159]) == 26
assert all(abs(x['val_tx_acc'] - 100/6) < 1e-10 for x in x600 if x['epoch'] >= 159)
assert len([x for x in x600 if x['epoch'] >= 149]) == 36
assert all(x['train_ecrs_cross_rx_identity_grad_norm'] == 0 for x in x600 if x['epoch'] >= 149)
assert x600[72]['train_ecrs_cross_rx_identity_grad_norm'] is None
assert x600[-1]['train_grad_backbone'] > 0
wb = load_workbook(out / 'supporting_data.xlsx', read_only=True)
counts = {'All 58 rows': 58, 'Activation metrics': 5374, 'ECRS effects': 8,
          'Extended curves': sum(len(x) for x in curves.values())}
for name, count in counts.items():
    assert wb[name].max_row == count + 1
wb.close()

dest = Path('E:/type10-7/automation_reports/CV-SincNet/mechanism_activation_audit_20260910')
dest.mkdir(parents=True, exist_ok=True)
files = [x for x in out.iterdir() if x.is_file()]
for x in files:
    shutil.copyfile(x, dest / x.name)
    assert x.read_bytes() == (dest / x.name).read_bytes()
for folder in [out, dest]:
    for link in re.findall(r'\]\(([^)]+)\)', report):
        if link.startswith('http'):
            continue
        path = Path(link) if re.match(r'^[A-Z]:', link) else folder / link
        assert path.exists(), path
print(json.dumps({'status': 'VERIFIED', 'sheets': counts, 'matched_ecrs_rows': 8,
                  'source_chance_streak': 26, 'zero_probe_streak': 36,
                  'mirrored_files': len(files)}, ensure_ascii=False))
