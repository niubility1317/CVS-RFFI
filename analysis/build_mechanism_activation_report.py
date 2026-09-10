"""Build a frozen, read-only mechanism audit from complete-log summaries."""
import csv
import gzip
import json
import shutil
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent
OUT = BASE / 'mechanism_activation_audit_20260910'
D = json.load(gzip.open(OUT / 'evidence.json.gz', 'rt', encoding='utf-8'))
CURVES = json.loads((OUT / 'extended_signal_curves.json').read_text(encoding='utf-8'))
SCENES = ['leo_clear_weak', 'leo_low_elev_weak', 'leo_rain_weak']

def options(x):
    argv = x['process'].get('argv', [])
    result = {}
    for i, v in enumerate(argv):
        if v.startswith('--'):
            result[v[2:]] = argv[i + 1] if i + 1 < len(argv) and not argv[i + 1].startswith('--') else True
    return result

def stat(x, field, prop='last'):
    return x['stats'].get(field, {}).get(prop)

def csvfile(name, rows):
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)

def state(run, name, x):
    if 'mechanism_screen' in run:
        return 'USER_RECONFIGURED_STOPPED' if x['jsonl_records'] or x['process'] else 'REPLACED_BEFORE_START'
    if not x['jsonl_records'] and not x['live_argv'] and not x['errors']:
        return 'RESERVED_NO_EPOCH'
    if 'selected_adv3b02' in run:
        return 'HISTORICAL_TARGET_CONTAMINATED'
    if 'matched_core90' in run:
        return 'HISTORICAL_INCOMPLETE'
    if x['errors']:
        return 'TRAIN_FAILED'
    if run.startswith('a1_extended') and name == 'X2_E600':
        return 'RUNNING_FUNCTIONALLY_COLLAPSED'
    if x['live_argv']:
        return 'RUNNING_AT_SNAPSHOT'
    return x['state'].get('status', 'UNKNOWN')

inventory, metrics = [], []
for run, r in D['runs'].items():
    for name, x in r['rows'].items():
        assert not x['parse_errors'] and not x['csv_bad_widths']
        assert x['jsonl_records'] == x['csv_records']
        o = options(x)
        inventory.append(dict(run=run, row=name, assessed_state=state(run, name, x),
            stored_state=x['state'].get('status'), epochs=x['jsonl_records'], live=bool(x['live_argv']),
            target_score_files=len(x['scores']), final_checkpoint=x['final_checkpoint_exists'],
            cross_rx_weight=o.get('a1_ecrs_cross_rx_weight'),
            cross_rx_executed_epochs=stat(x, 'train_ecrs_cross_rx_executed', 'nonzero_epochs'),
            cross_rx_gradient_nonzero_epochs=stat(x, 'train_ecrs_cross_rx_identity_grad_norm', 'nonzero_epochs'),
            cross_rx_leo_first=stat(x, 'train_ecrs_cross_rx_leo_count', 'first_nonzero'),
            eta_valid_nonzero_epochs=stat(x, 'train_r3_l_s_eta_valid', 'nonzero_epochs'),
            response_rho=o.get('a1_response_rho'), response_ce_epochs=stat(x, 'train_response_ce_weighted', 'nonzero_epochs'),
            response_pair_first=stat(x, 'train_response_pair_weighted', 'first_nonzero'),
            balanced_sampler=o.get('use_tx_rx_balanced_sampler'),
            fisher_weighted_epochs=stat(x, 'train_fisher_loss_weighted', 'nonzero_epochs'),
            last_source_clean=stat(x, 'val_tx_acc'), last_source_leo=stat(x, 'stage_source_val_sat_mean_tx'),
            errors=' | '.join(x['errors']), actual_argv=json.dumps(x['process'].get('argv', []), ensure_ascii=False)))
        for k, v in x['stats'].items():
            metrics.append(dict(run=run, row=name, metric=k, **v))
csvfile('all_rows.csv', inventory)
csvfile('activation_metrics.csv', metrics)

xrows = D['runs']['a1_ecrs_cross_rx_s392005_20260909_r1']['rows']
effects = []
for n, x in xrows.items():
    m = x['scores'][-1]['score']['metrics']
    assert all(m[s]['total'] == 168000 for s in ['clean'] + SCENES)
    effects.append(dict(row=n, target_clean=100*m['clean']['accuracy'],
        target_leo_mean=sum(100*m[s]['accuracy'] for s in SCENES)/3,
        **{'target_class_'+str(c):sum(100*m[s]['per_class_accuracy'][str(c)] for s in SCENES)/3 for c in range(6)}))
for x in effects:
    x['leo_delta_vs_x0_pp'] = x['target_leo_mean'] - effects[0]['target_leo_mean']
    x['leo_delta_vs_x2_pp'] = x['target_leo_mean'] - effects[2]['target_leo_mean']
csvfile('ecrs_effects.csv', effects)
curve_rows = [dict(row=n, **x) for n, rows in CURVES.items() for x in rows]
csvfile('extended_curves.csv', curve_rows)

plt.rcParams.update({'font.family':'DejaVu Sans', 'font.size':11, 'axes.spines.top':False, 'axes.spines.right':False})
fig, axes = plt.subplots(2, 1, figsize=(11, 7.2), sharex=True, constrained_layout=True)
for n, color in [('X2_E400','#2563eb'), ('X2_E600','#dc2626'), ('R3_CLEAN_RX_E400','#059669')]:
    rows = CURVES[n]
    axes[0].plot([x['epoch'] for x in rows], [x['val_tx_acc'] for x in rows], label=n, color=color, lw=1.7)
axes[0].axhline(100/6, color='#64748b', ls='--', lw=1, label='Six-class chance (16.67%)')
axes[0].set(ylabel='Source clean accuracy (%)', ylim=(10, 102), title='X2_E600: execution continues after source performance collapses')
axes[0].legend(loc='lower left', bbox_to_anchor=(0,1.05), ncol=2, frameon=False, fontsize=9)
rows = CURVES['X2_E600']
values = [x['train_ecrs_cross_rx_identity_grad_norm'] for x in rows]
axes[1].plot([x['epoch'] for x in rows], [float('nan') if v is None else v for v in values], color='#dc2626', lw=1.6)
axes[1].axvspan(149, rows[-1]['epoch'], color='#fecaca', alpha=.45)
axes[1].set_yscale('symlog', linthresh=1e-6)
axes[1].set(ylim=(-2e-7, .01), xlabel='Actual epoch (different training budgets)', ylabel='ECRS identity gradient probe', title='First-batch probe only; missing values are gaps, not zero')
axes[1].text(151, .0003, 'E149–E184: probe = 0', color='#991b1b')
for ax in axes:
    ax.grid(alpha=.16)
fig.savefig(OUT / 'extended_collapse.png', dpi=160)
plt.close(fig)

wb = Workbook(); wb.remove(wb.active)
for title, rows in [('All 58 rows', inventory), ('Activation metrics', metrics), ('ECRS effects', effects), ('Extended curves', curve_rows)]:
    ws = wb.create_sheet(title)
    ws.append(list(rows[0]))
    for r in rows:
        ws.append(list(r.values()))
    ws.freeze_panes = 'C2'; ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(color='FFFFFF', bold=True); cell.fill = PatternFill('solid', fgColor='17365D')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(55, max(14, len(str(col[0].value))+2))
wb.save(OUT / 'supporting_data.xlsx')

table = '\n'.join(f"|{x['row']}|{x['target_clean']:.4f}|{x['target_leo_mean']:.4f}|{x['leo_delta_vs_x0_pp']:+.4f}|{x['leo_delta_vs_x2_pp']:+.4f}|" for x in effects)
tot_epochs = sum(x['epochs'] for x in inventory)
tot_lines = sum(x['stdout_lines'] for r in D['runs'].values() for x in r['rows'].values())
assert len(inventory) == 58 and tot_epochs == 6884 and tot_lines == 310961
p = D['runs']['a1_mechanism_periodic_s392005_20260910_r1']['rows']
def final_leo(x):
    m = x['scores'][-1]['score']['metrics']
    return sum(m[s]['accuracy']*100 for s in SCENES)/3
f0, f3 = final_leo(p['F0_R3_REFERENCE']), final_leo(p['F3_R3_SWAP'])
report = f'''# 对话内实验的机制启用与效果审计

截至2026-09-10北京时间22:20的完整主快照；X2_E600曲线补读至E184。范围为本对话13个A1相关run根目录、58条实验/历史/占位记录。完整解析{tot_epochs}条epoch记录并扫描{tot_lines}行训练stdout，JSONL与CSV行数一致、未见解析或列宽错误。58条不等于58个完成实验，活动run也不是全服务器任务清单。

## 结论

有ECRS实验，而且X1—X7的跨接收机身份损失确实执行并产生非零梯度。但它们检验的是ECRS判别机制在A1的适配，不是完整物理响应ECRS。X1相对匹配X0仅提高target LEO均值0.3921个百分点；增加LEO配对没有继续提升，均衡采样和组合版下降明显。单seed不能认定稳定有效。

确实存在“方法名称覆盖的内容多于实际启用内容”：R3没有η监督和完整循环重编码，X系列没有物理响应估计/复杂锚点/融合。另有按设计关闭的对照、尚未到日程的机制、启动后数值失败的实验。不能把它们合写成“都没有启用”，也不能只凭非零loss宣布机制有效。

最严重的活动异常是X2_E600：源域clean从E70附近98.44%下降至六类机会水平；E159—E184连续26轮为16.6667%，E149—E184连续36轮ECRS身份梯度探针为0。进程运行与有效学习已明显分离。本次仅只读分析，未修改、停止或重启实验。

## ECRS的真实启用与匹配效果

X1—X7均完成E200和四场景独立评分，每场景168000条记录。`configured`、`executed`、raw/weighted loss及身份梯度探针在各自200轮中均非零；X1每batch平均127.916个合法anchor，接近128，并非长期没有合法配对。X2—X7的真实LEO参与从E80开始，121轮非零；X1全程LEO count=0符合clean-only设计。

`train_ecrs_cross_rx_successful_steps`是累计计数经过batch均值归约后的日志量，会出现44269.5这样的值。本报告只把它用作成功更新存在的佐证，不将其当成精确整数更新次数。梯度字段是每epoch首batch探针，不能代表逐batch全覆盖。

|实验|target clean/%|target LEO均值/%|LEO相对X0/pp|LEO相对X2/pp|
|---|---:|---:|---:|---:|
{table}

解释应按匹配关系进行：X1−X0检验clean跨RX；X2−X1检验追加LEO视图，差值−0.3093pp；X3—X6各自与X2比较，X7是多项组合不能单因归因。X1类1只提高约0.39pp，类3略降，尚未解决弱类问题。X1与另一矩阵的B1有多个设置差异，不能用二者差分证明某个机制效果。

X3的warm EMA实际衰减最小值约0.97105，X2约0.99；最终都接近0.999不意味着启动平均没用。X4/X7的coverage标志全程为1，DAOT目标进入日程后才具有对应蒸馏含义。X5/X7的`BATCH-GEOM`明确记录均衡采样启用，配对覆盖充分；X5的下降不能解释成“采样器忘了打开”。它改变了TX/RX/day组合和重复轨迹，可能影响身份学习分布，原因需source分层统计验证。

本矩阵设计文档的EX7明确为deferred：物理估计、复杂锚点和固定融合本轮不实现。因此正确名称是“A1+ECRS跨RX身份约束”。[设计依据](../../docs/A1_ECRS_CROSS_RX_ADAPTATION_20260909.md)

## 名称、实际行为与缺口逐项区分

|对象|实际证据|判断|
|---|---|---|
|X1—X7跨RX|200轮都有weighted loss与非零身份梯度探针；X2—X7从E80有真实LEO样本|已启用；效果有限或负向|
|X0跨RX|权重关闭，successful-steps全0|匹配消融，按设计关闭|
|X系列完整ECRS物理分支|EX7明确deferred|未实现；不能称完整ECRS复现|
|R3的η元数据监督|所有有该字段的source R3行，L_s/U_s eta_valid、eta raw/weighted都为0|实际未监督；元数据未接入|
|R3完整循环/重编码与其他配对|reencode两个参数均None；content/fingerprint配对标为invalid|当前适配未包含完整FCR机制|
|R3 self/swap/shared|非结构对照有相应非零加权损失，swap依日程进入|不能因η为0否定整个R3已启用部分|
|R3_STRUCTURE_CONTROL|aux_scale=0；保留相同前向、BN/RNG路径|结构对照，辅助损失有意关闭|
|E0_RESPONSE_ONLY|rho=0；response CE在124轮非零，有效数128|响应支路训练；响应到身份决策的固定融合有意关闭|
|E1_RESPONSE_FUSED|rho=0.05；42轮响应CE非零，E43触发非有限批次保护|支路已训练；正式日志缺直接融合增量/投影梯度证据；训练失败|
|E2_RESPONSE_PAIR|54轮响应CE非零；pair weighted从E21起34轮非零；E55失败|配对已启用后失败，不能写从未启用|
|F2_R3_IDENTITY|actual argv启用identity coupling；self/swap weighted非零|R3训练存在；缺正式coupling专用梯度/身份增量日志，不提升为端到端效果证实|
|G0_EQUAL_BRANCH|E1第8batch保护中止，无完整epoch|失败；无有效完整训练结论，也不推断一次forward都未执行|
|G1_FISHER_GATE|33轮branch CE/weighted loss非零，source evidence更新；null均值0.9555→0.6083|已启用且非全null；尚无target有效性结论|
|D0_NO_ORBIT|DAOT全程0，LEO监督CE仍从E80启用|只消融orbit，不是所有卫星机制关闭|
|截图B0/B1/B2/D0/D1/D2/D3/F3的cross-RX|successful-steps全0；B1 balanced_sampler=0|这八行没有追加X系列跨RX，不能外推到全部对话实验|
|早期mechanism_screen同名行|无活进程，有用户重配置停止记录；部分pipeline_state仍写RUNNING|被后续periodic替换；旧状态不可当运行事实|

R3代码在[objective第71行](../../code/cvsrffi/a1_r3_objective.py)创建零η、全false有效掩码，第85行传入空重编码；模块开头已说明η元数据不可用。这是有记录的范围裁剪，当前证据不支持把它称为新发现的静默开关故障。若设计声明要求完整FCR，则这些就是实质实现缺口。独立响应估计器的固定/无梯度求解与后端响应编码器可训练是不同层，不应混为“响应完全没有训练”。

当前证据也没有证明E1/F2的命名路径一定不影响身份决策。缺失探针应记为UNKNOWN，而不是把未观察到当作0。G1总weighted loss非零不等于其中每个分项每轮都已加权，route KL原始值与其日程权重需要分开理解。

## 运行中延长实验：日程与失效不能混淆

|行|完整epoch|跨RX状态|LEO参与|target结论|
|---|---:|---|---|---|
|X2_E400|173|所有已记录轮次有非零身份梯度探针|从实际E159进入，已有15轮|尚未到首次目标评估点|
|X2_E600|184，主快照183|早期启用，后期身份探针连续0；源域崩塌|截至该快照尚未到伸缩日程起点|没有target改善证据；当前学习状态异常|
|R3_CLEAN_RX_E400|114|clean跨RX已执行且梯度非零|cross-RX的LEO count=0符合clean-only；不代表R3自己的LEO重建关闭|尚未完成|
|R3_REFERENCE_CLEAN_CROSS_RX E200|200|clean跨RX200轮非零梯度|cross-RX仅clean|source训练完成，尚无target评分|

延长预算采用reference epoch映射，不能拿E200版本的实际E80门槛直接套到E400/E600。X2_E600的LEO count=0本身属于未到日程；其源域准确率崩塌则是另一个已观察到的问题。

![源域崩塌与梯度探针](extended_collapse.png)

X2_E600在E70的source clean为98.4370%，E73为77.8444%，E74为55.9741%，E75为19.6259%。E84首次精确落在16.6667%，后续长期处于机会水平，最近连续段为E159—E184。E184的source三LEO均值也是16.6667%。

同时，E184的合法anchor仍约127.87，executed=1，raw triplet loss约0.200000046，等于margin=0.2，加权后约0.010000002；成功计数继续增加，但身份梯度探针为0。这与退化身份表示下triplet项不再提供有效区分信号相容，说明“合法样本+非零loss+成功step”不足以证明机制在工作。

限制必须保留：E73等若干轮probe缺失是null，不能当成0；这里的0只覆盖首batch探针，不证明所有batch或全模型梯度为0。E184主干总梯度约0.322仍非零，TX CE约12.02。未见本行Traceback/OOM或系统性保护终止，因此不是进程未启动。现有证据尚不能把根因唯一归为ECRS、长日程、域对抗或某一次更新。

下一步应围绕E68—E85核查身份特征范数/方差、类别logit塌缩、TX与域/辅助梯度关系及optimizer数值状态，先定位退化起点，再决定修复。该项属于后续技术诊断建议；本报告不以低分为理由擅停活动run。

## 当前16行periodic矩阵的完成与失败

主快照已有B0/B1/B2/D0/D1/D2/D3/F0/F3共9行E200评分；此前方法相关性报告冻结的八行没有F0，本报告保留其原样本集合，不把后来完成的F0悄悄混入旧相关系数。F1=E192、F2=E161、E0=E124、G1=E33仍在运行。E1在E43第38batch报8次非有限，E2在E55第88batch报8次，G0在E1前8batch全部非有限。

新增完成的F0_R3_REFERENCE target LEO均值为{f0:.6f}%，F3_R3_SWAP为{f3:.6f}%，匹配差值只有{f3-f0:+.6f}pp。即使F3相对B0或别的行改善，也不能把跨骨架差异全归因于swap。F0/F3这组才更接近对应机制比较。

其他历史记录：A1 fast V2四行和R3 scratch三行已有E200评分；R3 B120/B160第二次发布已各自完成预算及评分，首次占位没有epoch。两个早期CORE90匹配根目录没有实验行，第三次仅29轮历史未完成记录。继承ADV3B02旧checkpoint的两行保留历史结果，但因已有目标接触不进入干净泛化证据。所有这些状态都可在58行总表定位。

## 对增强路线的修正

现在优先补source的class×RX×scene失败分布和梯度/身份几何诊断。X1对X0是小幅正向证据，值得保留为匹配参考，但不足以直接升级为默认方案；X2追加LEO并未改善，不能假定“更多视图一定更好”。

均衡采样已经有X5/X7负向证据，应先复盘覆盖、样本重复和类间margin，避免原样再跑。类×域尾部风险与带可靠性约束的U覆盖仍是待验证方向，但只有source诊断明确支持时再设计小范围单因素实验。先把现有分支的身份贡献测清楚，再考虑补全昂贵物理分支，能够减少“名义增强、实际无贡献”的解释空间。

协议保持source-only选择、checkpoint全继承来源检查与truth-last评分。上述探索性target结果已被用于形成研发假设，不能继续在同一目标记录上宣称独立确认或单seed晋级。数据角色和物理样本不变不新增重复数据验证。

## 可复核产物

- [58行状态与actual argv](all_rows.csv)、[所有机制字段的全程统计](activation_metrics.csv)、[ECRS匹配效果](ecrs_effects.csv)、[延长实验完整曲线](extended_curves.csv)。
- [可筛选工作簿](supporting_data.xlsx)、[完整解析证据摘要](evidence.json.gz)、[曲线补读原始JSON](extended_signal_curves.json)。证据包保存全日志解析后的统计、字段名、配置、错误、评分和进程快照，不声称包含全部原始stdout文本。
- [只读采集脚本](../audit_conversation_mechanisms.py)、[完整曲线采集](../inspect_extended_signal.py)、[本报告计算与构建](../build_mechanism_activation_report.py)。
- [不同方法与逐类跨域分析](../method_class_analysis_20260910/report.md)。

状态结论：数据和上述机制事实VERIFIED；E1/E2/G0训练FAILED；E1融合专用贡献、F2耦合专用梯度及X2_E600唯一根因UNKNOWN。运行状态只对应标注快照时刻。
'''
for relative in ['../../docs/A1_ECRS_CROSS_RX_ADAPTATION_20260909.md', '../../code/cvsrffi/a1_r3_objective.py', '../audit_conversation_mechanisms.py', '../inspect_extended_signal.py', '../build_mechanism_activation_report.py']:
    report = report.replace('](' + relative + ')', '](' + (OUT / relative).resolve().as_posix() + ')')
(OUT / 'report.md').write_text(report, encoding='utf-8')
print(json.dumps({'rows':len(inventory), 'epochs':tot_epochs, 'stdout_lines':tot_lines, 'metric_rows':len(metrics), 'f3_minus_f0':f3-f0, 'report_chars':len(report)}, ensure_ascii=False))
