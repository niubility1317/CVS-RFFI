"""Export verified CSV tables to an editable workbook and check row counts."""
import argparse,csv
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment

def main():
    p=argparse.ArgumentParser();p.add_argument('folder',type=Path);a=p.parse_args()
    wb=Workbook();wb.remove(wb.active);counts={}
    order=['summary','seed_summary','paired_seed_summary','DR_EG_interaction','paired_rescue_harm','all_groups','coverage36','source_curves','independent_recount','artifact_paths']
    for path in sorted(a.folder.glob('*.csv'),key=lambda p:order.index(p.stem) if p.stem in order else len(order)):
        with path.open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
        ws=wb.create_sheet(path.stem[:31]);ws.append(rows[0])
        for values in rows[1:]:
            converted=[]
            for header,value in zip(rows[0],values):
                try:
                    if header in ('row','row_id','method','candidate','reference','axis','key','scene','seeds','confusion','per_class_f1'):raise ValueError()
                    number=float(value);converted.append(int(number) if number.is_integer() else number)
                except ValueError:converted.append(value)
            ws.append(converted)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='163B5C')
        for i,header in enumerate(rows[0],1):
            ws.column_dimensions[ws.cell(1,i).column_letter].width=24 if header not in ('confusion','checkpoint','predictions','score','source_log') else 55
            if (any(x in header for x in ('accuracy','macro_f1','worst_RX','worst_TX')) and not header.endswith('_pp')) or (path.stem=='seed_summary' and header in ('mean','sd')):
                for col in ws.iter_cols(min_col=i,max_col=i,min_row=2):
                    for cell in col:cell.number_format='0.0000%'
        counts[ws.title]=len(rows)
    output=a.folder/'complete_test_data.xlsx';wb.save(output)
    check=load_workbook(output,read_only=True,data_only=True)
    assert {s.title:s.max_row for s in check}==counts
    check.close();print('WORKBOOK_VERIFIED '+str(counts))

if __name__=='__main__':main()
